"""Formatting layer: sanitize AI text and export it as HTML preview, DOCX and PDF."""
import html
import re
from io import BytesIO

import config
from ai_core.utils import split_terms

# --------------------------------------------------------------------------- sanitize
_CHAR_MAP = {
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u2013": "-", "\u2014": " - ", "\u2212": "-",
    "\u2026": "...", "\u00a0": " ", "\u2022": "-", "\u25cf": "-", "\u25aa": "-",
    "\u200b": "", "\u200c": "", "\u200d": "", "\ufeff": "",
}


def sanitize_text(text: str) -> str:
    """Remove typographic quotes, markdown symbols and stray control characters."""
    if not text:
        return ""
    for src, dst in _CHAR_MAP.items():
        text = text.replace(src, dst)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"^[ \t]{0,3}#{1,6}[ \t]*", "", text, flags=re.M)      # markdown headings
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)                          # bold
    text = re.sub(r"^[ \t]*[*]\s+", "- ", text, flags=re.M)               # "* item" bullets
    text = text.replace("`", "")
    text = re.sub(r"^[ \t]*-{3,}[ \t]*$", "", text, flags=re.M)           # horizontal rules
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- parsing
_BULLET = re.compile(r"^(?:[-\u2022]\s+|\(?[a-zA-Z0-9]{1,2}\)\s+)(.+)$")
_NUM_HEADING = re.compile(r"^\d+(?:\.\d+)*[.)]?\s+\S")


def _is_heading(line: str) -> bool:
    s = line.strip()
    if len(s) > 90:
        return False
    if s.endswith(":"):
        return True
    if _NUM_HEADING.match(s) and not s.endswith((".", ";", ",")):
        return True
    core = s.rstrip(":").strip()
    return len(core) > 3 and core.upper() == core and any(c.isalpha() for c in core)


def parse_document(text: str, doc_type: str = "") -> list:
    """Return a list of (kind, text) with kind in title | heading | para | bullet."""
    lines = [ln.strip() for ln in sanitize_text(text).split("\n") if ln.strip()]
    blocks = []
    for i, line in enumerate(lines):
        if i == 0 and len(line) <= 100 and not line.endswith((".", ";", ":")):
            blocks.append(("title", line))
        elif _BULLET.match(line) and not _NUM_HEADING.match(line):
            blocks.append(("bullet", line))
        elif _is_heading(line):
            blocks.append(("heading", line))
        else:
            blocks.append(("para", line))
    if not blocks or blocks[0][0] != "title":
        blocks.insert(0, ("title", doc_type.strip() or "Legal Document"))
    return blocks


def _split_title(blocks: list):
    return blocks[0][1], blocks[1:]


# --------------------------------------------------------------------------- HTML preview
def format_html_preview(text: str) -> str:
    """Stylised HTML (no blank lines, so Streamlit's markdown leaves it alone)."""

    def esc(s: str) -> str:
        return html.escape(s).replace("$", "&#36;")  # '$' would trigger LaTeX in st.markdown

    parts, in_list = [], False
    for kind, content in parse_document(text):
        if kind != "bullet" and in_list:
            parts.append("</ul>")
            in_list = False
        if kind == "title":
            parts.append(
                f"<h2 style='text-align:center;margin:0 0 16px 0;color:#f3f4f6;'>{esc(content)}</h2>"
            )
        elif kind == "heading":
            parts.append(f"<h4 style='margin:18px 0 6px 0;color:#93c5fd;'>{esc(content)}</h4>")
        elif kind == "bullet":
            if not in_list:
                parts.append("<ul style='margin:4px 0 4px 18px;padding:0;'>")
                in_list = True
            parts.append(f"<li style='margin:3px 0;'>{esc(_BULLET.match(content).group(1))}</li>")
        else:
            parts.append(f"<p style='margin:6px 0;line-height:1.6;color:#e5e7eb;'>{esc(content)}</p>")
    if in_list:
        parts.append("</ul>")
    return "".join(parts)


# --------------------------------------------------------------------------- DOCX
def _add_field(run, instruction: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def format_docx(text: str, doc_type: str, terms: str = None) -> bytes:
    """Word document: logo, Times New Roman, title, terms table, body, footer with page numbers."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    config.ensure_logos()
    title, blocks = _split_title(parse_document(text, doc_type))

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(config.LOGO_PATH), width=Inches(2.2))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    term_list = split_terms(terms or "")
    if term_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.add_run("Key Terms Summary").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "No.", "Term / Condition"
        for cell in hdr:
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, term in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text, row[1].text = str(i), term
        for row in table.rows:
            row.cells[0].width = Inches(0.6)
            row.cells[1].width = Inches(5.9)
        doc.add_paragraph()

    for kind, content in blocks:
        if kind == "heading":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.keep_with_next = True
            para.add_run(content).bold = True
        elif kind == "bullet":
            doc.add_paragraph(_BULLET.match(content).group(1), style="List Bullet")
        else:
            para = doc.add_paragraph(content)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(f"{config.FOOTER_TEXT}  |  Page ")
    fr.italic = True
    fr.font.size = Pt(9)
    num = footer.add_run()
    num.italic = True
    num.font.size = Pt(9)
    _add_field(num, "PAGE")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF
def _latin1(s: str) -> str:
    """fpdf2 core fonts only support latin-1."""
    return s.encode("latin-1", "replace").decode("latin-1")


def format_pdf(text: str, doc_type: str, terms: str = None) -> bytes:
    """Branded PDF: logo + title header and footer on every page, bold headings, bullet terms."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    from PIL import Image

    config.ensure_logos()
    title, blocks = _split_title(parse_document(text, doc_type))
    logo_path = str(config.LOGO_PATH)
    with Image.open(logo_path) as im:
        ratio = im.height / im.width

    class LegalPDF(FPDF):
        def header(self):
            w = 42
            self.image(logo_path, x=(self.w - w) / 2, y=8, w=w)
            self.set_y(8 + w * ratio + 2)
            self.set_font("Helvetica", "B", 12)
            self.set_text_color(20, 30, 60)
            self.cell(0, 7, _latin1(title), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.ln(4)
            self.set_text_color(0, 0, 0)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(110, 110, 110)
            self.cell(0, 8, _latin1(f"{config.FOOTER_TEXT}  |  Page {self.page_no()}/{{nb}}"), align="C")
            self.set_text_color(0, 0, 0)

    pdf = LegalPDF(format="A4")
    pdf.alias_nb_pages()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    def heading(txt):
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 6, _latin1(txt), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(1)

    def para(txt):
        pdf.set_font("Times", "", 11)
        pdf.multi_cell(0, 5.8, _latin1(txt), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(1.5)

    def bullet(txt):
        pdf.set_font("Times", "", 11)
        pdf.set_x(pdf.l_margin + 4)
        pdf.cell(5, 5.8, "-")
        pdf.multi_cell(0, 5.8, _latin1(txt), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(0.8)

    term_list = split_terms(terms or "")
    if term_list:
        heading("Key Terms Summary")
        for term in term_list:
            bullet(term)
        pdf.ln(3)

    for kind, content in blocks:
        if kind == "heading":
            heading(content)
        elif kind == "bullet":
            bullet(_BULLET.match(content).group(1))
        else:
            para(content)

    return bytes(pdf.output())
